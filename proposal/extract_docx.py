"""Temporary script to extract .docx content to text."""
from docx import Document
import os

doc = Document('Proposal_2211102226_Revisi 2 .docx')

output = []

for para in doc.paragraphs:
    style = para.style.name if para.style else ''
    text = para.text.strip()
    if not text:
        output.append('')
        continue
    
    # Mark headings
    if 'Heading' in style:
        level = style.replace('Heading ', '').replace('Heading', '1')
        try:
            level = int(level)
        except:
            level = 1
        output.append(f"{'#' * level} {text}")
    else:
        output.append(text)

# Also extract tables
for i, table in enumerate(doc.tables):
    output.append(f"\n--- TABLE {i+1} ---")
    for row in table.rows:
        cells = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
        output.append(' | '.join(cells))
    output.append("--- END TABLE ---\n")

result = '\n'.join(output)
out_path = 'proposal_extracted.txt'
with open(out_path, 'w', encoding='utf-8') as f:
    f.write(result)

print(f"Extracted {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
print(f"Saved to {out_path} ({len(result)} chars)")
